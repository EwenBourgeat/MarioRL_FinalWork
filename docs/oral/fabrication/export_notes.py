"""Écrit docs/oral/Notes_oral.md : le texte à dire, diapo par diapo, lu dans
les notes de l'orateur du .pptx (copie de secours si Canva ne les importe pas).

    python docs/oral/fabrication/export_notes.py
"""

import re
from pathlib import Path

from pptx import Presentation

ORAL = Path(__file__).resolve().parents[1]
PARTS = {1: "Introduction", 3: "Monde 1-1 : la veille", 9: "Monde 1-2 : le duel sur v3",
         14: "Monde 1-3 : le vrai jeu NES", 18: "Monde 1-4 : le recul", 22: "Fin", 24: "Annexes (seulement si le jury pose une question)"}


def title_of(slide):
    for shape in slide.shapes:
        if shape.has_text_frame and shape.name in ("titre", "grand_titre") and shape.text_frame.text.strip():
            t = " ".join(shape.text_frame.text.split())
            return t.capitalize() if t.isupper() else t
    return "-"


def to_md(notes):
    out, duration = [], None
    for line in notes.splitlines():
        line = line.strip()
        m = re.fullmatch(r"\[(\d+:\d\d)\]", line)
        if m:
            duration = m.group(1)
            continue
        for label in ("À l'écran :", "On dit :", "Transition :"):
            if line.startswith(label):
                rest = line[len(label):].strip()
                line = f"**{label[:-2]}** : {rest}" if rest else f"**{label[:-2]}** :"
                if label == "On dit :":
                    out.append("")
                if label == "Transition :":
                    out.append("")
                    line = f"> ➜ {rest}"
                break
        out.append(line)
    return duration, "\n".join(out).strip()


def main():
    prs = Presentation(ORAL / "Oral_RL_Mario.pptx")
    body, total = [], 0
    for i, slide in enumerate(prs.slides, 1):
        if i in PARTS:
            body += [f"## {PARTS[i]}", ""]
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        duration, md = to_md(notes)
        if duration:
            m, s = map(int, duration.split(":"))
            total += 60 * m + s
        head = f"### Diapo {i} · {title_of(slide)}" + (f"  ({duration})" if duration else "")
        body += [head, "", md or "-", ""]
    intro = ["# Oral final : ce qu'on dit", "",
             "Pour chaque diapo : ce qu'on voit à l'écran, les phrases à dire, puis la phrase pour passer à la suivante.",
             f"Durée prévue : {total // 60} min {total % 60:02d} s (entre parenthèses, le temps de chaque diapo).",
             "", "Conseils : parler lentement, montrer l'écran du doigt, et lancer les vidéos quand c'est indiqué.", ""]
    (ORAL / "Notes_oral.md").write_text("\n".join(intro + body), encoding="utf-8")
    print("Notes_oral.md écrit,", len(prs.slides), "diapos,", total // 60, "min", total % 60, "s")


if __name__ == "__main__":
    main()
